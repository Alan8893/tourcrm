"""HTTP-level integration tests for Participant Export (TH-0118.4 /
Issue #218; docs/05-api/participant-export-api.md, PO decisions GAP-1..7
recorded on Issue #218):

    GET  /api/v1/memberships/exports/fields
    POST /api/v1/memberships/exports

Against the REAL shipped app (app.main.app) and a real PostgreSQL database,
with the seeded canonical `admin` Role and its migration-seeded grants —
mirroring tests/integration/test_membership_imports_api.py's fixture
conventions (local, duplicated helpers).
"""

import io
import re
import uuid
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from app.api.deps import CurrentPrincipal, get_current_principal
from app.db.authorization import (
    Permission,
    Role,
    RolePermission,
    RolePermissionScope,
    UserRoleAssignment,
)
from app.db.events import Event, EventParticipation
from app.db.groups import Group, GroupMembership
from app.db.identity import Club, ClubMembership, GuardianRelationship, Person, User
from app.db.session import session_scope
from app.main import app

from .conftest import requires_postgres

pytestmark = requires_postgres

_URL = "/api/v1/memberships/exports"
_FIELDS_URL = "/api/v1/memberships/exports/fields"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_READ_PERMISSIONS = (
    "person.read",
    "membership.read",
    "group.read",
    "event.read",
    "guardian_relationship.read",
)


@pytest.fixture
def client() -> TestClient:
    test_client = TestClient(app, raise_server_exceptions=True)
    yield test_client
    app.dependency_overrides.clear()


def _utc(*args: int) -> datetime:
    return datetime(*args, tzinfo=timezone.utc)


# --- fixtures / factories ---------------------------------------------------


def _make_club(session) -> Club:
    club = Club(name=f"Club {uuid.uuid4().hex[:8]}", status="active")
    session.add(club)
    session.flush()
    return club


def _make_person(session, last_name: str, first_name: str, **overrides: object) -> Person:
    person = Person(last_name=last_name, first_name=first_name, **overrides)
    session.add(person)
    session.flush()
    return person


def _make_club_membership(session, club: Club, person: Person, **overrides) -> ClubMembership:
    defaults: dict[str, object] = {
        "club_id": club.id,
        "person_id": person.id,
        "membership_type": "member",
        "status": "active",
        "joined_at": _utc(2020, 1, 1),
    }
    defaults.update(overrides)
    membership = ClubMembership(**defaults)  # type: ignore[arg-type]
    session.add(membership)
    session.flush()
    return membership


def _make_group(session, club: Club, name: str) -> Group:
    group = Group(club_id=club.id, name=name, status="active", valid_from=_utc(2020, 1, 1))
    session.add(group)
    session.flush()
    return group


def _make_group_membership(session, group: Group, membership: ClubMembership, **overrides):
    defaults: dict[str, object] = {
        "group_id": group.id,
        "club_membership_id": membership.id,
        "valid_from": _utc(2024, 1, 1),
        "membership_status": "active",
    }
    defaults.update(overrides)
    session.add(GroupMembership(**defaults))  # type: ignore[arg-type]
    session.flush()


def _make_event(session, club: Club, title: str) -> Event:
    event = Event(
        club_id=club.id,
        event_type="trip",
        title=title,
        start_at=_utc(2026, 10, 1, 10, 0),
        end_at=_utc(2026, 10, 1, 18, 0),
        timezone="Europe/Moscow",
        status="published",
    )
    session.add(event)
    session.flush()
    return event


def _participate(session, event: Event, person: Person, status: str = "registered") -> None:
    session.add(
        EventParticipation(event_id=event.id, person_id=person.id, registration_status=status)
    )
    session.flush()


def _guardian(session, guardian: Person, child: Person, **overrides) -> None:
    defaults: dict[str, object] = {
        "guardian_person_id": guardian.id,
        "child_person_id": child.id,
        "relationship_type": "parent",
        "status": "active",
        "valid_from": _utc(2020, 1, 1),
    }
    defaults.update(overrides)
    session.add(GuardianRelationship(**defaults))  # type: ignore[arg-type]
    session.flush()


def _make_user(session, person: Person | None = None) -> User:
    if person is None:
        person = _make_person(session, "Админова", f"U-{uuid.uuid4().hex[:6]}")
    user = User(
        person=person,
        login_identifier=f"user-{uuid.uuid4().hex[:8]}@example.com",
        status="active",
    )
    session.add(user)
    session.flush()
    return user


def _assign_system_role(
    user_id: uuid.UUID,
    role_code: str,
    *,
    club_id: uuid.UUID | None = None,
    valid_to: datetime | None = None,
) -> None:
    with session_scope() as session:
        role = session.execute(select(Role).where(Role.code == role_code)).scalar_one()
        kwargs: dict[str, object] = {}
        if valid_to is not None:
            kwargs["valid_from"] = valid_to - timedelta(days=30)
            kwargs["valid_to"] = valid_to
        session.add(
            UserRoleAssignment(
                user_id=user_id, role_id=role.id, scope_type="all", club_id=club_id, **kwargs
            )
        )
        session.commit()


def _grant_permission(user_id: uuid.UUID, permission_code: str, club_id: uuid.UUID) -> None:
    """Ad hoc `all` grant through a throwaway, non-system role."""
    with session_scope() as session:
        permission = session.execute(
            select(Permission).where(Permission.code == permission_code)
        ).scalar_one()
        role = Role(code=f"role-{uuid.uuid4().hex[:8]}", name="Test role")
        session.add(role)
        session.flush()
        session.add(
            RolePermission(
                role_id=role.id,
                permission_id=permission.id,
                scopes=[RolePermissionScope(scope_type="all")],
            )
        )
        session.add(UserRoleAssignment(user_id=user_id, role_id=role.id, club_id=club_id))
        session.commit()


def _revoke_admin_grant(permission_code: str) -> None:
    with session_scope() as session:
        admin_role_id = session.execute(select(Role.id).where(Role.code == "admin")).scalar_one()
        permission_id = session.execute(
            select(Permission.id).where(Permission.code == permission_code)
        ).scalar_one()
        session.execute(
            delete(RolePermission).where(
                RolePermission.role_id == admin_role_id,
                RolePermission.permission_id == permission_id,
            )
        )
        session.commit()


def _authenticate_as(user_id: uuid.UUID) -> None:
    app.dependency_overrides[get_current_principal] = lambda: CurrentPrincipal(
        user_id=user_id, session_id=uuid.uuid4()
    )


def _csrf_headers(client: TestClient) -> dict:
    client.cookies.set("csrf_token", "test-csrf-token")
    return {"X-CSRF-Token": "test-csrf-token"}


class Scenario:
    club_id: uuid.UUID
    group_id: uuid.UUID
    event_id: uuid.UUID
    admin_user_id: uuid.UUID


def _scenario() -> Scenario:
    """One Club with:

    - Алексеева Анна — active CM; active GM in the Group; registered for
      the Event; guardians Петров Пётр Петрович (two active relationships
      of different types, one phone) and Сидорова Ольга (no phone), plus a
      revoked and an expired relationship that must not appear;
    - Борисов Борис — active CM; an ended and a later active GM in the
      Group (historical duplicate); not participating;
    - Васильева Вера — active CM; no Group; `declined` for the Event;
    - Григорьев Глеб — suspended CM; active GM in the Group; registered;
    - Дмитриева Дарья — two active CMs of different types (duplicate);
    - Егоров Егор — ended GM only; no Event.
    """
    scenario = Scenario()
    with session_scope() as session:
        club = _make_club(session)
        group = _make_group(session, club, "Юные туристы")
        event = _make_event(session, club, "Осенний поход")

        anna = _make_person(
            session,
            "Алексеева",
            "Анна",
            middle_name="Игоревна",
            birth_date=date(2012, 3, 4),
            phone="+7 900 000-00-01",
            email="anna@example.com",
            address="г. Москва",
        )
        anna_cm = _make_club_membership(session, club, anna)
        _make_group_membership(session, group, anna_cm)
        _participate(session, event, anna)
        petrov = _make_person(session, "Петров", "Пётр", middle_name="Петрович", phone="+71")
        sidorova = _make_person(session, "Сидорова", "Ольга")
        revoked = _make_person(session, "Отозванный", "Опекун", phone="+79")
        expired = _make_person(session, "Истёкший", "Опекун", phone="+78")
        _guardian(session, sidorova, anna)
        _guardian(session, petrov, anna)
        _guardian(session, petrov, anna, relationship_type="legal_representative")
        _guardian(session, revoked, anna, status="revoked", relationship_type="other")
        _guardian(
            session, expired, anna, valid_from=_utc(2019, 1, 1), valid_to=_utc(2020, 1, 1)
        )

        boris = _make_person(session, "Борисов", "Борис")
        boris_cm = _make_club_membership(session, club, boris)
        _make_group_membership(
            session,
            group,
            boris_cm,
            valid_from=_utc(2022, 1, 1),
            valid_to=_utc(2023, 1, 1),
            membership_status="ended",
        )
        _make_group_membership(session, group, boris_cm, valid_from=_utc(2024, 1, 1))

        vera = _make_person(session, "Васильева", "Вера")
        _make_club_membership(session, club, vera)
        _participate(session, event, vera, status="declined")

        gleb = _make_person(session, "Григорьев", "Глеб")
        gleb_cm = _make_club_membership(session, club, gleb, status="suspended")
        _make_group_membership(session, group, gleb_cm)
        _participate(session, event, gleb)

        darya = _make_person(session, "Дмитриева", "Дарья")
        _make_club_membership(session, club, darya)
        _make_club_membership(session, club, darya, membership_type="sport")

        egor = _make_person(session, "Егоров", "Егор")
        egor_cm = _make_club_membership(session, club, egor)
        _make_group_membership(
            session,
            group,
            egor_cm,
            valid_from=_utc(2022, 1, 1),
            valid_to=_utc(2023, 1, 1),
            membership_status="ended",
        )

        admin = _make_user(session)
        session.commit()
        scenario.club_id = club.id
        scenario.group_id = group.id
        scenario.event_id = event.id
        scenario.admin_user_id = admin.id
    _assign_system_role(scenario.admin_user_id, "admin")
    return scenario


def _export(client: TestClient, **body):
    body.setdefault("format", "print")
    body.setdefault("fields", ["person.last_name", "person.first_name"])
    return client.post(_URL, json=body, headers=_csrf_headers(client))


def _print_rows(response) -> list[list[str]]:
    assert response.status_code == 200, response.text
    tbody = re.search(r"<tbody>(.*)</tbody>", response.text, re.S)
    assert tbody is not None
    return [
        re.findall(r"<td>(.*?)</td>", row) for row in re.findall(r"<tr>(.*?)</tr>", tbody.group(1))
    ]


def _last_names(response) -> list[str]:
    return [row[0] for row in _print_rows(response)]


@pytest.fixture
def scenario(client: TestClient) -> Scenario:
    data = _scenario()
    _authenticate_as(data.admin_user_id)
    return data


# --- contexts / dataset semantics -------------------------------------------


def test_club_context_defaults_to_active_club_membership_one_row_per_person(
    client, scenario
) -> None:
    response = _export(client, context="club")
    # Григорьев (suspended) excluded; Дмитриева (two active CMs) once; the
    # admin's own Person has no ClubMembership in this Club.
    assert _last_names(response) == ["Алексеева", "Борисов", "Васильева", "Дмитриева", "Егоров"]


def test_club_context_membership_status_filter(client, scenario) -> None:
    response = _export(
        client,
        context="club",
        membership_status="suspended",
        fields=["person.last_name", "membership.status"],
    )
    assert _print_rows(response) == [["Григорьев", "suspended"]]


def test_group_context_defaults_to_active_group_membership(client, scenario) -> None:
    response = _export(
        client,
        context="group",
        group_id=str(scenario.group_id),
        fields=["person.last_name", "group.name", "membership.status"],
    )
    # Борисов has an ended and an active GM: one row. Егоров (ended only)
    # is excluded by the `active` default.
    assert _print_rows(response) == [
        ["Алексеева", "Юные туристы", "active"],
        ["Борисов", "Юные туристы", "active"],
        ["Григорьев", "Юные туристы", "active"],
    ]


def test_group_context_membership_status_ended(client, scenario) -> None:
    response = _export(
        client, context="group", group_id=str(scenario.group_id), membership_status="ended"
    )
    assert _last_names(response) == ["Борисов", "Егоров"]


def test_event_context_all_participation_statuses_with_active_club_membership(
    client, scenario
) -> None:
    response = _export(
        client,
        context="event",
        event_id=str(scenario.event_id),
        fields=[
            "person.last_name",
            "event.name",
            "event.starts_at",
            "event_participation.status",
            "membership.status",
        ],
    )
    # Starts 10:00 UTC → 13:00 in the Event's timezone (Europe/Moscow).
    assert _print_rows(response) == [
        ["Алексеева", "Осенний поход", "2026-10-01 13:00", "registered", "active"],
        ["Васильева", "Осенний поход", "2026-10-01 13:00", "declined", "active"],
    ]


def test_event_context_participation_status_filter(client, scenario) -> None:
    response = _export(
        client,
        context="event",
        event_id=str(scenario.event_id),
        participation_status="registered",
    )
    assert _last_names(response) == ["Алексеева"]


def test_group_event_context_is_group_membership_and_participation_intersection(
    client, scenario
) -> None:
    response = _export(
        client,
        context="group_event",
        group_id=str(scenario.group_id),
        event_id=str(scenario.event_id),
        fields=["person.last_name", "group.name", "event.name", "event_participation.status"],
    )
    # Борисов: group only; Васильева: event only — both excluded.
    assert _print_rows(response) == [
        ["Алексеева", "Юные туристы", "Осенний поход", "registered"],
        ["Григорьев", "Юные туристы", "Осенний поход", "registered"],
    ]


def test_empty_result_set(client, scenario) -> None:
    response = _export(
        client,
        context="group_event",
        group_id=str(scenario.group_id),
        event_id=str(scenario.event_id),
        participation_status="declined",
    )
    assert _print_rows(response) == []
    assert "Участники не найдены." in response.text


# --- field selection ---------------------------------------------------------


def test_only_selected_fields_in_requested_order(client, scenario) -> None:
    response = _export(
        client,
        context="club",
        membership_status="active",
        fields=[
            "person.email",
            "person.last_name",
            "person.birth_date",
            "person.middle_name",
            "person.phone",
            "person.address",
        ],
    )
    headers = re.findall(r"<th>(.*?)</th>", response.text)
    assert headers == ["Email", "Фамилия", "Дата рождения", "Отчество", "Телефон", "Адрес"]
    assert _print_rows(response)[0] == [
        "anna@example.com",
        "Алексеева",
        "2012-03-04",
        "Игоревна",
        "+7 900 000-00-01",
        "г. Москва",
    ]


def test_guardian_fields_join_active_guardians_in_stable_order(client, scenario) -> None:
    response = _export(
        client, context="club", fields=["person.last_name", "guardian.name", "guardian.phone"]
    )
    rows = {row[0]: row[1:] for row in _print_rows(response)}
    assert rows["Алексеева"] == ["Петров Пётр Петрович; Сидорова Ольга", "+71"]
    assert rows["Борисов"] == ["", ""]


def test_guardian_fields_are_never_included_implicitly(client, scenario) -> None:
    response = _export(client, context="club")
    assert "Петров" not in response.text
    assert "Представитель" not in response.text


# --- output formats (one dataset) --------------------------------------------


def test_xlsx_output(client, scenario) -> None:
    response = _export(
        client,
        context="event",
        event_id=str(scenario.event_id),
        format="xlsx",
        fields=["person.last_name", "person.birth_date", "event.starts_at"],
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == _XLSX_MIME
    assert response.headers["content-disposition"].startswith("attachment;")
    assert ".xlsx" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == ["Фамилия", "Дата рождения", "Начало мероприятия"]
    assert [c.value for c in sheet[2]] == [
        "Алексеева",
        datetime(2012, 3, 4),
        datetime(2026, 10, 1, 13, 0),
    ]
    assert sheet.max_row == 3


def test_pdf_output(client, scenario) -> None:
    response = _export(client, context="club", format="pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert ".pdf" in response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF-")


def test_print_output_is_self_contained_html(client, scenario) -> None:
    response = _export(client, context="club", format="print")
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/html; charset=utf-8"
    assert "content-disposition" not in response.headers
    assert response.headers["content-security-policy"].startswith("default-src 'none'")
    assert "@media print" in response.text
    assert "<script" not in response.text


def test_all_formats_render_the_same_dataset(client, scenario) -> None:
    body = {
        "context": "group",
        "group_id": str(scenario.group_id),
        "fields": ["person.last_name", "person.first_name", "group.name"],
    }
    printed = _print_rows(_export(client, format="print", **body))
    sheet = load_workbook(io.BytesIO(_export(client, format="xlsx", **body).content)).active
    assert [[c.value for c in row] for row in sheet.iter_rows(min_row=2)] == printed
    pdf = _export(client, format="pdf", **body)
    assert pdf.status_code == 200


def test_export_changes_nothing(client, scenario) -> None:
    def counts() -> tuple[int, ...]:
        with session_scope() as session:
            return tuple(
                session.execute(select(func.count()).select_from(model)).scalar_one()
                for model in (Person, ClubMembership, GroupMembership, EventParticipation)
            )

    before = counts()
    for fmt in ("xlsx", "pdf", "print"):
        assert _export(client, context="club", format=fmt).status_code == 200
    assert counts() == before


# --- authorization -----------------------------------------------------------


def test_unauthenticated_is_401(client) -> None:
    response = client.post(
        _URL,
        json={"context": "club", "fields": ["person.last_name"], "format": "xlsx"},
        headers=_csrf_headers(client),
    )
    assert response.status_code == 401
    assert client.get(_FIELDS_URL).status_code == 401


def test_missing_csrf_token_is_rejected(client, scenario) -> None:
    response = client.post(
        _URL, json={"context": "club", "fields": ["person.last_name"], "format": "xlsx"}
    )
    assert response.status_code == 403


def _user_without_admin_role(club_id: uuid.UUID) -> uuid.UUID:
    with session_scope() as session:
        user = _make_user(session)
        session.commit()
        return user.id


@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "print"])
def test_non_administrator_with_all_read_grants_is_denied(client, scenario, fmt) -> None:
    user_id = _user_without_admin_role(scenario.club_id)
    for code in _READ_PERMISSIONS:
        _grant_permission(user_id, code, scenario.club_id)
    _authenticate_as(user_id)
    response = _export(client, context="club", format=fmt)
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "forbidden"
    assert client.get(_FIELDS_URL).status_code == 403


@pytest.mark.parametrize("role_code", ["instructor", "member", "guardian"])
def test_other_system_roles_are_denied(client, scenario, role_code) -> None:
    user_id = _user_without_admin_role(scenario.club_id)
    _assign_system_role(user_id, role_code, club_id=scenario.club_id)
    _authenticate_as(user_id)
    assert _export(client, context="club").status_code == 403


def test_expired_admin_assignment_is_denied(client, scenario) -> None:
    user_id = _user_without_admin_role(scenario.club_id)
    _assign_system_role(user_id, "admin", valid_to=datetime.now(timezone.utc) - timedelta(days=1))
    _authenticate_as(user_id)
    assert _export(client, context="club").status_code == 403


def test_club_scoped_admin_of_current_club_is_allowed(client, scenario) -> None:
    user_id = _user_without_admin_role(scenario.club_id)
    _assign_system_role(user_id, "admin", club_id=scenario.club_id)
    _authenticate_as(user_id)
    assert _export(client, context="club").status_code == 200


def test_admin_scoped_to_another_club_is_denied(client, scenario, monkeypatch) -> None:
    with session_scope() as session:
        other_club_id = _make_club(session).id
        session.commit()
    # Two Clubs violate the single-Club product invariant, so the current
    # Club is pinned the way resolve_sole_club_id would resolve it.
    monkeypatch.setattr("app.exports.service.resolve_sole_club_id", lambda _s: scenario.club_id)
    user_id = _user_without_admin_role(scenario.club_id)
    _assign_system_role(user_id, "admin", club_id=other_club_id)
    _authenticate_as(user_id)
    assert _export(client, context="club").status_code == 403


def test_field_requires_its_existing_read_grant(client, scenario) -> None:
    """Export never widens rights: without `guardian_relationship.read`
    the admin cannot export guardian fields, but person fields still work."""
    _revoke_admin_grant("guardian_relationship.read")
    assert _export(client, context="club").status_code == 200
    response = _export(client, context="club", fields=["person.last_name", "guardian.name"])
    assert response.status_code == 403


def test_context_requires_its_existing_read_grant(client, scenario) -> None:
    _revoke_admin_grant("event.read")
    assert _export(client, context="club").status_code == 200
    response = _export(client, context="event", event_id=str(scenario.event_id))
    assert response.status_code == 403


def test_non_all_scope_read_grant_is_not_sufficient(client, scenario) -> None:
    """An instructor-style `own_groups` person.read does not satisfy the
    club-wide read requirement, even together with the admin role."""
    _revoke_admin_grant("person.read")
    with session_scope() as session:
        permission = session.execute(
            select(Permission).where(Permission.code == "person.read")
        ).scalar_one()
        role = Role(code=f"role-{uuid.uuid4().hex[:8]}", name="Test role")
        session.add(role)
        session.flush()
        session.add(
            RolePermission(
                role_id=role.id,
                permission_id=permission.id,
                scopes=[RolePermissionScope(scope_type="own_groups")],
            )
        )
        session.add(
            UserRoleAssignment(
                user_id=scenario.admin_user_id, role_id=role.id, club_id=scenario.club_id
            )
        )
        session.commit()
    assert _export(client, context="club").status_code == 403


# --- validation / existence hiding ------------------------------------------


@pytest.mark.parametrize(
    ("fields", "code"),
    [
        ([], "empty_export_fields"),
        (["person.last_name", "person.id"], "unknown_export_field"),
        (["user.password_hash"], "unknown_export_field"),
        (["person.photo_file_id"], "unknown_export_field"),
        (["document.medical_certificate"], "unknown_export_field"),
        (["person.last_name", "group.name"], "export_field_not_available"),
        (["person.last_name", "person.last_name"], "duplicate_export_field"),
    ],
)
def test_invalid_fields_are_rejected(client, scenario, fields, code) -> None:
    response = _export(client, context="club", fields=fields)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == code


@pytest.mark.parametrize(
    ("body", "code"),
    [
        ({"context": "group"}, "export_context_target_required"),
        ({"context": "event"}, "export_context_target_required"),
        ({"context": "club", "participation_status": "registered"}, "export_filter_not_applicable"),
        ({"context": "club", "membership_status": "ended"}, "invalid_membership_status"),
    ],
)
def test_invalid_context_filter_combinations_are_rejected(client, scenario, body, code) -> None:
    response = _export(client, **body)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == code


def test_group_id_in_event_context_is_rejected(client, scenario) -> None:
    response = _export(
        client, context="event", event_id=str(scenario.event_id), group_id=str(scenario.group_id)
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "export_filter_not_applicable"


@pytest.mark.parametrize(
    "body",
    [
        {"context": "club", "person_ids": []},
        {"context": "people"},
        {"context": "club", "format": "csv"},
    ],
)
def test_malformed_request_is_rejected(client, scenario, body) -> None:
    response = _export(client, **body)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_validation_is_not_reachable_without_administrator(client, scenario) -> None:
    user_id = _user_without_admin_role(scenario.club_id)
    _authenticate_as(user_id)
    assert _export(client, context="club", fields=["user.password_hash"]).status_code == 403


def test_nonexistent_group_and_event_are_404(client, scenario) -> None:
    group = _export(client, context="group", group_id=str(uuid.uuid4()))
    event = _export(client, context="event", event_id=str(uuid.uuid4()))
    assert group.status_code == 404 and event.status_code == 404
    assert group.json()["error"]["code"] == "not_found"


def test_group_and_event_of_another_club_are_indistinguishable_from_nonexistent(
    client, scenario, monkeypatch
) -> None:
    with session_scope() as session:
        other = _make_club(session)
        other_group_id = _make_group(session, other, "Чужая группа").id
        other_event_id = _make_event(session, other, "Чужой поход").id
        session.commit()
    monkeypatch.setattr("app.exports.service.resolve_sole_club_id", lambda _s: scenario.club_id)

    def body(response) -> dict:
        error = response.json()["error"]
        return {"code": error["code"], "message": error["message"]}

    foreign_group = _export(client, context="group", group_id=str(other_group_id))
    missing_group = _export(client, context="group", group_id=str(uuid.uuid4()))
    assert foreign_group.status_code == missing_group.status_code == 404
    assert body(foreign_group) == body(missing_group)

    foreign_event = _export(client, context="event", event_id=str(other_event_id))
    missing_event = _export(client, context="event", event_id=str(uuid.uuid4()))
    assert foreign_event.status_code == missing_event.status_code == 404
    assert body(foreign_event) == body(missing_event)


# --- fields endpoint ---------------------------------------------------------


def test_fields_endpoint_returns_canonical_allowlist(client, scenario) -> None:
    response = client.get(_FIELDS_URL)
    assert response.status_code == 200
    items = response.json()["items"]
    assert [item["field_code"] for item in items] == [
        "person.last_name",
        "person.first_name",
        "person.middle_name",
        "person.birth_date",
        "person.phone",
        "person.email",
        "person.address",
        "group.name",
        "membership.status",
        "event.name",
        "event.starts_at",
        "event_participation.status",
        "guardian.name",
        "guardian.phone",
    ]
    by_code = {item["field_code"]: item for item in items}
    assert by_code["group.name"]["contexts"] == ["group", "group_event"]
    assert by_code["event.starts_at"]["contexts"] == ["event", "group_event"]
    assert by_code["person.last_name"]["label"] == "Фамилия"
