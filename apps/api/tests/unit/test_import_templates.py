"""Downloadable participant import templates (Issue #228).

The committed static templates (apps/web/public/templates/) must carry
exactly the canonical import columns of app.imports.parsing, in contract
order, with no data rows, and must be accepted by the existing parser.
Regenerate with `python -m app.cli.generate_import_templates`.
"""

import io
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.cli.generate_import_templates import TEMPLATES_DIR
from app.imports.parsing import (
    CANONICAL_IMPORT_COLUMNS,
    OPTIONAL_IMPORT_COLUMNS,
    REQUIRED_IMPORT_COLUMNS,
    parse_source,
)
from app.imports.templates import (
    CSV_TEMPLATE_FILENAME,
    TEMPLATE_COLUMNS,
    XLSX_TEMPLATE_FILENAME,
    build_csv_template,
    build_xlsx_template,
)

CSV_PATH: Path = TEMPLATES_DIR / CSV_TEMPLATE_FILENAME
XLSX_PATH: Path = TEMPLATES_DIR / XLSX_TEMPLATE_FILENAME
REGENERATE = "run `python -m app.cli.generate_import_templates` from apps/api"


def test_template_columns_are_the_canonical_contract_in_order() -> None:
    assert TEMPLATE_COLUMNS == REQUIRED_IMPORT_COLUMNS + OPTIONAL_IMPORT_COLUMNS
    assert set(TEMPLATE_COLUMNS) == CANONICAL_IMPORT_COLUMNS
    assert len(TEMPLATE_COLUMNS) == len(CANONICAL_IMPORT_COLUMNS)


def test_committed_csv_template_matches_the_generator() -> None:
    assert CSV_PATH.read_bytes() == build_csv_template(), REGENERATE


def test_committed_xlsx_template_has_only_the_canonical_header() -> None:
    workbook = load_workbook(io.BytesIO(XLSX_PATH.read_bytes()), read_only=True)
    assert len(workbook.worksheets) == 1, REGENERATE
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    assert rows == [TEMPLATE_COLUMNS], REGENERATE


def test_generated_xlsx_has_only_the_canonical_header() -> None:
    workbook = load_workbook(io.BytesIO(build_xlsx_template()), read_only=True)
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    assert rows == [TEMPLATE_COLUMNS]


@pytest.mark.parametrize(("path", "source_format"), [(CSV_PATH, "csv"), (XLSX_PATH, "xlsx")])
def test_existing_parser_accepts_the_empty_template(path: Path, source_format: str) -> None:
    parsed = parse_source(path.read_bytes(), source_format)
    assert parsed.columns == TEMPLATE_COLUMNS
    assert parsed.records == ()


def test_filled_csv_template_parses_as_records() -> None:
    row = "Иван,Петров,,2010-05-01,,ivan@example.com,\r\n"
    filled = CSV_PATH.read_bytes() + row.encode("utf-8")
    parsed = parse_source(filled, "csv")
    assert parsed.columns == TEMPLATE_COLUMNS
    assert len(parsed.records) == 1
    assert parsed.records[0].values["first_name"] == "Иван"
    assert parsed.records[0].values["last_name"] == "Петров"


def test_filled_xlsx_template_parses_as_records() -> None:
    workbook = load_workbook(io.BytesIO(XLSX_PATH.read_bytes()))
    sheet = workbook.worksheets[0]
    sheet.append(["Иван", "Петров", None, date(2010, 5, 1), None, "ivan@example.com", None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parse_source(buffer.getvalue(), "xlsx")
    assert parsed.columns == TEMPLATE_COLUMNS
    assert len(parsed.records) == 1
    assert parsed.records[0].values["last_name"] == "Петров"
