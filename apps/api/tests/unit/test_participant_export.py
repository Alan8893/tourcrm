"""Unit tests for Participant Export (TH-0118.4 / Issue #218) — the
canonical field allowlist, request validation and the three renderers.
No HTTP, no database."""

import io
import uuid
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.exports.fields import (
    EXPORT_FIELDS,
    EXPORT_FIELDS_BY_CODE,
    context_permissions,
    field_permission,
    membership_status_source,
)
from app.exports.filters import (
    CANONICAL_PARTICIPATION_STATUSES,
    PARTICIPATION_STATUS_OPTIONS,
)
from app.exports.queries import GuardianContact
from app.exports.rendering import (
    format_cell,
    render_pdf,
    render_print_html,
    render_xlsx,
)
from app.exports.service import (
    ExportRequestError,
    ParticipantExportDataset,
    ParticipantExportRequest,
    guardian_cells,
    validate_export_request,
)

_GROUP_ID = uuid.uuid4()
_EVENT_ID = uuid.uuid4()

# participant-export-api.md §5 — the complete MVP allowlist, nothing more.
_CANONICAL_FIELD_CODES = {
    "person.last_name",
    "person.first_name",
    "person.middle_name",
    "person.birth_date",
    "person.phone",
    "person.email",
    "person.address",
    "group.name",
    "membership.status",
    "event.name",
    "event.starts_at",
    "event_participation.status",
    "guardian.name",
    "guardian.phone",
}


def _request(context: str, fields: tuple[str, ...] = ("person.last_name",), **kwargs):
    ids = {}
    if context in ("group", "group_event"):
        ids["group_id"] = _GROUP_ID
    if context in ("event", "group_event"):
        ids["event_id"] = _EVENT_ID
    ids.update(kwargs)
    return ParticipantExportRequest(context=context, fields=fields, **ids)


# --- canonical allowlist ----------------------------------------------------


def test_allowlist_is_exactly_the_canonical_field_set() -> None:
    assert {f.code for f in EXPORT_FIELDS} == _CANONICAL_FIELD_CODES
    assert len(EXPORT_FIELDS) == len(_CANONICAL_FIELD_CODES)


@pytest.mark.parametrize(
    "code",
    [
        "person.id",
        "person.photo_file_id",
        "user.password_hash",
        "user.login_identifier",
        "session.token",
        "document.file",
        "document.storage_key",
        "medical_certificate",
        "person.medical_notes",
    ],
)
def test_sensitive_and_internal_fields_are_not_in_allowlist(code: str) -> None:
    assert code not in EXPORT_FIELDS_BY_CODE


def test_every_field_has_label_and_permission_in_every_available_context() -> None:
    for field in EXPORT_FIELDS:
        assert field.label
        for context in field.contexts:
            assert field_permission(field.code, context).endswith(".read")


def test_field_availability_per_context() -> None:
    assert EXPORT_FIELDS_BY_CODE["group.name"].contexts == {"group", "group_event"}
    for code in ("event.name", "event.starts_at", "event_participation.status"):
        assert EXPORT_FIELDS_BY_CODE[code].contexts == {"event", "group_event"}
    for code in ("person.last_name", "membership.status", "guardian.name", "guardian.phone"):
        assert EXPORT_FIELDS_BY_CODE[code].contexts == {"club", "group", "event", "group_event"}


def test_membership_status_source_per_context() -> None:
    assert membership_status_source("club") == "club_membership"
    assert membership_status_source("event") == "club_membership"
    assert membership_status_source("group") == "group_membership"
    assert membership_status_source("group_event") == "group_membership"


def test_permission_mapping_uses_only_existing_read_permissions() -> None:
    assert context_permissions("club") == {"person.read", "membership.read"}
    assert context_permissions("group") == {"person.read", "group.read"}
    assert context_permissions("event") == {"person.read", "membership.read", "event.read"}
    assert context_permissions("group_event") == {"person.read", "group.read", "event.read"}
    assert field_permission("guardian.name", "club") == "guardian_relationship.read"
    assert field_permission("membership.status", "club") == "membership.read"
    assert field_permission("membership.status", "group") == "group.read"
    assert field_permission("event_participation.status", "event") == "event.read"


# --- request validation -----------------------------------------------------


@pytest.mark.parametrize("context", ["club", "group", "event", "group_event"])
def test_valid_request_defaults_membership_status_to_active(context: str) -> None:
    status, columns = validate_export_request(_request(context))
    assert status == "active"
    assert [c.code for c in columns] == ["person.last_name"]


def test_columns_keep_requested_order() -> None:
    _, columns = validate_export_request(
        _request("club", ("person.first_name", "person.last_name", "guardian.name"))
    )
    assert [c.code for c in columns] == ["person.first_name", "person.last_name", "guardian.name"]


def _error(request: ParticipantExportRequest) -> ExportRequestError:
    with pytest.raises(ExportRequestError) as info:
        validate_export_request(request)
    return info.value


def test_empty_fields_rejected() -> None:
    assert _error(_request("club", ())).code == "empty_export_fields"


def test_unknown_field_rejected() -> None:
    error = _error(_request("club", ("person.last_name", "user.password_hash")))
    assert error.code == "unknown_export_field"
    assert error.details["fields"] == ["user.password_hash"]


def test_duplicate_field_rejected() -> None:
    error = _error(_request("club", ("person.last_name", "person.last_name")))
    assert error.code == "duplicate_export_field"


@pytest.mark.parametrize(
    ("context", "code"),
    [
        ("club", "group.name"),
        ("club", "event.name"),
        ("group", "event.starts_at"),
        ("group", "event_participation.status"),
        ("event", "group.name"),
    ],
)
def test_field_not_available_in_context_rejected(context: str, code: str) -> None:
    error = _error(_request(context, ("person.last_name", code)))
    assert error.code == "export_field_not_available"
    assert error.details["fields"] == [code]


@pytest.mark.parametrize(
    ("context", "missing"),
    [("group", "group_id"), ("event", "event_id"), ("group_event", "group_id")],
)
def test_missing_context_target_rejected(context: str, missing: str) -> None:
    error = _error(_request(context, **{missing: None}))
    assert error.code == "export_context_target_required"
    assert error.details["field"] == missing


@pytest.mark.parametrize(
    ("context", "extra"),
    [
        ("club", {"group_id": _GROUP_ID}),
        ("club", {"event_id": _EVENT_ID}),
        ("club", {"participation_status": "registered"}),
        ("group", {"event_id": _EVENT_ID}),
        ("group", {"participation_status": "registered"}),
        ("event", {"group_id": _GROUP_ID}),
    ],
)
def test_filter_not_applicable_to_context_rejected(context: str, extra: dict) -> None:
    error = _error(_request(context, **extra))
    assert error.code == "export_filter_not_applicable"
    assert set(error.details["fields"]) == set(extra)


@pytest.mark.parametrize(
    ("context", "status"),
    [("club", "ended"), ("event", "ended"), ("group", "suspended"), ("group_event", "pending")],
)
def test_membership_status_vocabulary_follows_context_record(context: str, status: str) -> None:
    assert _error(_request(context, membership_status=status)).code == "invalid_membership_status"


@pytest.mark.parametrize(
    ("context", "status"),
    [("club", "suspended"), ("event", "archived"), ("group", "ended"), ("group_event", "ended")],
)
def test_membership_status_accepted_for_context_record(context: str, status: str) -> None:
    assert validate_export_request(_request(context, membership_status=status))[0] == status


# --- participation_status vocabulary (PR #226 PO decision) ----------------


def test_participation_status_vocabulary_is_the_implemented_mvp_statuses() -> None:
    # ADR-0037: the only statuses the self-registration workflow writes.
    from app.events.participation import CANCELLED_STATUS, REGISTERED_STATUS

    assert CANONICAL_PARTICIPATION_STATUSES == {REGISTERED_STATUS, CANCELLED_STATUS}
    assert [option.value for option in PARTICIPATION_STATUS_OPTIONS] == ["registered", "cancelled"]
    assert all(option.label for option in PARTICIPATION_STATUS_OPTIONS)


@pytest.mark.parametrize("context", ["event", "group_event"])
@pytest.mark.parametrize("status", ["registered", "cancelled"])
def test_canonical_participation_status_accepted(context: str, status: str) -> None:
    validate_export_request(_request(context, participation_status=status))


@pytest.mark.parametrize("context", ["event", "group_event"])
@pytest.mark.parametrize("status", ["invited", "waitlisted", "declined", "removed", "Registered"])
def test_unknown_participation_status_rejected(context: str, status: str) -> None:
    error = _error(_request(context, participation_status=status))
    assert error.code == "invalid_participation_status"
    assert error.details["allowed"] == ["cancelled", "registered"]


def test_participation_status_on_non_event_context_is_not_applicable_first() -> None:
    # Applicability is reported before the vocabulary check.
    error = _error(_request("club", participation_status="invited"))
    assert error.code == "export_filter_not_applicable"


def test_invalid_context_rejected() -> None:
    error = _error(ParticipantExportRequest(context="person_ids", fields=("person.last_name",)))
    assert error.code == "invalid_export_context"


# --- guardian cells ---------------------------------------------------------


def test_guardian_cells_keep_phone_positions_aligned_with_names() -> None:
    contacts = [
        GuardianContact(name="Андреева Алла", phone=None),
        GuardianContact(name="Петров Пётр Петрович", phone="+7 900 000-00-01"),
        GuardianContact(name="Сидорова Ольга", phone=""),
    ]
    names, phones = guardian_cells(contacts)
    assert names == "Андреева Алла; Петров Пётр Петрович; Сидорова Ольга"
    assert phones == "—; +7 900 000-00-01; —"


def test_guardian_cells_without_guardians_are_empty() -> None:
    assert guardian_cells([]) == (None, None)


# --- renderers ----------------------------------------------------------------


def _dataset(rows: tuple = ()) -> ParticipantExportDataset:
    columns = tuple(
        EXPORT_FIELDS_BY_CODE[code]
        for code in ("person.last_name", "person.birth_date", "event.starts_at", "person.address")
    )
    return ParticipantExportDataset(
        title="Участники мероприятия «Поход <b>»",
        generated_at=datetime(2026, 9, 30, 12, 0, tzinfo=timezone.utc),
        columns=columns,
        rows=rows,
    )


_ROW = (
    "Иванова",
    date(2012, 3, 4),
    datetime(2026, 10, 1, 13, 0, tzinfo=ZoneInfo("Europe/Moscow")),
    '=HYPERLINK("http://x")',
)


def test_format_cell() -> None:
    assert format_cell(None) == ""
    assert format_cell(date(2012, 3, 4)) == "2012-03-04"
    assert format_cell(datetime(2026, 10, 1, 13, 5)) == "2026-10-01 13:05"
    assert format_cell("a\x01b") == "ab"


def test_xlsx_has_header_rows_and_typed_cells() -> None:
    workbook = load_workbook(io.BytesIO(render_xlsx(_dataset((_ROW,)))))
    sheet = workbook.active
    assert [c.value for c in sheet[1]] == [
        "Фамилия",
        "Дата рождения",
        "Начало мероприятия",
        "Адрес",
    ]
    assert sheet["A2"].value == "Иванова"
    assert sheet["B2"].value == datetime(2012, 3, 4)
    # Rendered in the Event's own timezone, stored as local wall-clock time.
    assert sheet["C2"].value == datetime(2026, 10, 1, 13, 0)
    assert sheet.max_row == 2


def test_xlsx_never_stores_values_as_formulas() -> None:
    workbook = load_workbook(io.BytesIO(render_xlsx(_dataset((_ROW,)))))
    cell = workbook.active["D2"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://x")'


def test_xlsx_empty_result_has_only_header() -> None:
    sheet = load_workbook(io.BytesIO(render_xlsx(_dataset()))).active
    assert sheet.max_row == 1


def test_pdf_is_generated_with_cyrillic_font() -> None:
    content = render_pdf(_dataset((_ROW,) * 50))
    assert content.startswith(b"%PDF-")
    assert b"DejaVu" in content


def test_pdf_empty_result() -> None:
    assert render_pdf(_dataset()).startswith(b"%PDF-")


def test_print_html_is_self_contained_escaped_and_print_ready() -> None:
    page = render_print_html(_dataset((("<script>x</script>", None, None, "a & b"),)))
    assert page.startswith("<!DOCTYPE html>")
    assert "@media print" in page
    assert "<script" not in page
    assert "&lt;script&gt;x&lt;/script&gt;" in page
    assert "Поход &lt;b&gt;" in page
    assert "a &amp; b" in page
    assert "http" not in page  # no external resource


def test_print_html_empty_result() -> None:
    page = render_print_html(_dataset())
    assert "Участники не найдены." in page
    assert "<tbody></tbody>" in page
