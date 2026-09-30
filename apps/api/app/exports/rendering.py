"""Participant Export output formats (participant-export-api.md §7).

Presentation only: every function here takes the already-authorized
`ParticipantExportDataset` and never queries, filters or authorizes.

- XLSX — editable spreadsheet (openpyxl, already used by TH-0118.2);
- PDF — print-ready document (fpdf2 with the bundled DejaVu Sans font, so
  Cyrillic renders without any system font — GAP-5);
- print — a self-contained `text/html` page with `@media print` rules, no
  scripts and no external resources (GAP-6).
"""

import html
import io
from datetime import date, datetime
from pathlib import Path

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, TYPE_STRING
from openpyxl.styles import Font

from app.exports.service import CellValue, ParticipantExportDataset

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"
PRINT_MEDIA_TYPE = "text/html; charset=utf-8"

_FONT_PATH = Path(__file__).parent / "fonts" / "DejaVuSans.ttf"
_FONT_FAMILY = "DejaVuSans"
_PDF_CELL_PADDING = 1
_PDF_MAX_COLUMN_TEXT_WIDTH = 45
_PDF_MAX_WORD_WIDTH = 25
_EMPTY_RESULT_TEXT = "Участники не найдены."
_GENERATED_LABEL = "Сформировано"
_DATE_FORMAT = "%Y-%m-%d"
_DATETIME_FORMAT = "%Y-%m-%d %H:%M"
# A spreadsheet treats text starting with these as a formula; exported
# values are always stored as literal text (CSV/formula injection).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def format_cell(value: CellValue) -> str:
    """Text form shared by the PDF and print formats."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(_DATETIME_FORMAT)
    if isinstance(value, date):
        return value.strftime(_DATE_FORMAT)
    # Control characters are not printable and not representable in XLSX XML.
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _generated_text(dataset: ParticipantExportDataset) -> str:
    return f"{_GENERATED_LABEL}: {dataset.generated_at.strftime(_DATETIME_FORMAT)} UTC"


def render_xlsx(dataset: ParticipantExportDataset) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Участники"
    sheet.append([column.label for column in dataset.columns])
    for header_cell in sheet[1]:
        header_cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for row_index, row in enumerate(dataset.rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(value, datetime):
                # openpyxl cannot store tz-aware datetimes; the value is
                # already in the Event's own timezone (Issue #218 decision).
                cell.value = value.replace(tzinfo=None)
                cell.number_format = "YYYY-MM-DD HH:MM"
            elif isinstance(value, date):
                cell.value = value
                cell.number_format = "YYYY-MM-DD"
            elif isinstance(value, str):
                text = format_cell(value)
                cell.value = text
                if text.startswith(_FORMULA_PREFIXES):
                    cell.data_type = TYPE_STRING
            else:
                cell.value = value

    for column_index, column in enumerate(dataset.columns, start=1):
        width = max(
            [len(column.label)] + [len(format_cell(row[column_index - 1])) for row in dataset.rows]
        )
        letter = sheet.cell(row=1, column=column_index).column_letter
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 60)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_column_widths(pdf: FPDF, dataset: ParticipantExportDataset) -> tuple[float, ...]:
    """Column widths (mm) from the rendered text width: every column keeps
    at least its longest header word (and, up to a cap, its longest value
    word) unbroken; the remaining page width is
    shared in proportion to how much more each column's content needs (long
    free text such as addresses is capped so it wraps instead of squeezing
    every other column)."""
    padding = 2 * _PDF_CELL_PADDING + 1
    minimum: list[float] = []
    desired: list[float] = []
    for index, column in enumerate(dataset.columns):
        header_word = max(pdf.get_string_width(word) for word in column.label.split())
        values = [format_cell(row[index]) for row in dataset.rows]
        longest_value = max((pdf.get_string_width(value) for value in values), default=0.0)
        longest_word = max(
            (pdf.get_string_width(word) for value in values for word in value.split()),
            default=0.0,
        )
        minimum.append(max(header_word, min(longest_word, _PDF_MAX_WORD_WIDTH)) + padding)
        desired.append(max(header_word, min(longest_value, _PDF_MAX_COLUMN_TEXT_WIDTH)) + padding)
    available = pdf.epw
    if sum(desired) <= available:
        return tuple(desired)
    spare = available - sum(minimum)
    growth = sum(d - m for d, m in zip(desired, minimum))
    if spare <= 0 or growth <= 0:
        return tuple(minimum)
    return tuple(m + spare * (d - m) / growth for d, m in zip(desired, minimum))


def render_pdf(dataset: ParticipantExportDataset) -> bytes:
    pdf = FPDF(orientation="landscape", unit="mm", format="A4")
    pdf.set_creator("TourCRM")
    pdf.set_title(dataset.title)
    pdf.add_font(_FONT_FAMILY, fname=str(_FONT_PATH))
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.set_margins(left=12, top=12, right=12)
    pdf.add_page()

    pdf.set_font(_FONT_FAMILY, size=14)
    pdf.multi_cell(0, 8, dataset.title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(_FONT_FAMILY, size=8)
    pdf.cell(0, 6, _generated_text(dataset), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font(_FONT_FAMILY, size=8)
    with pdf.table(
        col_widths=_pdf_column_widths(pdf, dataset),
        text_align="LEFT",
        padding=_PDF_CELL_PADDING,
        headings_style=FontFace(fill_color=(230, 230, 230)),
        line_height=4,
        repeat_headings=1,
    ) as table:
        header = table.row()
        for column in dataset.columns:
            header.cell(column.label)
        for values in dataset.rows:
            row = table.row()
            for value in values:
                row.cell(format_cell(value))

    if not dataset.rows:
        pdf.ln(2)
        pdf.cell(0, 6, _EMPTY_RESULT_TEXT, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def render_print_html(dataset: ParticipantExportDataset) -> str:
    esc = html.escape
    header = "".join(f"<th>{esc(column.label)}</th>" for column in dataset.columns)
    body = "".join(
        "<tr>" + "".join(f"<td>{esc(format_cell(value))}</td>" for value in row) + "</tr>"
        for row in dataset.rows
    )
    empty = "" if dataset.rows else f'<p class="empty">{esc(_EMPTY_RESULT_TEXT)}</p>'
    return (
        "<!DOCTYPE html>\n"
        '<html lang="ru">\n<head>\n<meta charset="utf-8">\n'
        f"<title>{esc(dataset.title)}</title>\n"
        "<style>\n"
        "body{font-family:'DejaVu Sans',Arial,sans-serif;font-size:11pt;color:#000;"
        "background:#fff;margin:16px}\n"
        "h1{font-size:16pt;margin:0 0 4px}\n"
        ".generated{font-size:9pt;color:#444;margin:0 0 12px}\n"
        "table{border-collapse:collapse;width:100%}\n"
        "th,td{border:1px solid #666;padding:4px 6px;text-align:left;vertical-align:top}\n"
        "th{background:#e6e6e6}\n"
        "thead{display:table-header-group}\n"
        "tr{page-break-inside:avoid}\n"
        "@media print{\n"
        "@page{size:A4 landscape;margin:12mm}\n"
        "body{margin:0;font-size:9pt}\n"
        "th{-webkit-print-color-adjust:exact;print-color-adjust:exact}\n"
        "}\n"
        "</style>\n</head>\n<body>\n"
        f"<h1>{esc(dataset.title)}</h1>\n"
        f'<p class="generated">{esc(_generated_text(dataset))}</p>\n'
        f"<table>\n<thead><tr>{header}</tr></thead>\n<tbody>{body}</tbody>\n</table>\n"
        f"{empty}\n"
        "</body>\n</html>\n"
    )


__all__ = [
    "PDF_MEDIA_TYPE",
    "PRINT_MEDIA_TYPE",
    "XLSX_MEDIA_TYPE",
    "format_cell",
    "render_pdf",
    "render_print_html",
    "render_xlsx",
]
