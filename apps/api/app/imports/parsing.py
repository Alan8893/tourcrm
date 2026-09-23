"""CSV/XLSX source-file parsing for participant import (TH-0118.2).

Canonical sources: docs/05-api/people-api.md §22 "Source file format",
docs/04-modules/people-and-membership.md §11.1.

Pure Python — no FastAPI, no database session, no FileStorage: takes the
source file's bytes and returns its header and records, or raises
`ImportParseError` for a file that cannot be parsed at all (which fails
the whole job). Row-level problems are NOT decided here — values are
returned raw and judged later by app.imports.rows.

Rules (both formats):

- the first row is the header, made of canonical machine-readable column
  names only (`CANONICAL_IMPORT_COLUMNS`) — no aliases, no guessing;
  surrounding whitespace in a header cell is ignored;
- an unknown column (including a blank column name), a repeated column,
  or a missing required column (`first_name`/`last_name`) is a header error;
- every following row is a record; a row whose cells are all empty is not
  a record and is ignored;
- `row_number` is the 1-based row number in the source file (the header
  is row 1), so it matches what the user sees in a spreadsheet.

CSV: UTF-8 (a leading BOM is allowed), comma-delimited, standard quoting.
Undecodable bytes, malformed quoting/NUL bytes, or a record carrying a
non-empty value beyond the last header column make the file malformed.

XLSX: read with openpyxl in read-only mode (the source is never written);
only the first worksheet is read — any further worksheets are ignored.
Cell values are returned as openpyxl produces them (str/int/float/
datetime/...); typing is judged per field by app.imports.rows.
"""

import csv
import io
from dataclasses import dataclass
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

REQUIRED_IMPORT_COLUMNS: tuple[str, ...] = ("first_name", "last_name")
OPTIONAL_IMPORT_COLUMNS: tuple[str, ...] = (
    "middle_name",
    "birth_date",
    "phone",
    "email",
    "external_id",
)
CANONICAL_IMPORT_COLUMNS: frozenset[str] = frozenset(
    REQUIRED_IMPORT_COLUMNS + OPTIONAL_IMPORT_COLUMNS
)

_HEADER_ROW_NUMBER = 1
_MAX_REPORTED_COLUMN_LENGTH = 255


@dataclass(frozen=True)
class SourceRecord:
    """One record: its source row number and raw value per header column."""

    row_number: int
    values: dict[str, Any]


@dataclass(frozen=True)
class ParsedSource:
    columns: tuple[str, ...]
    records: tuple[SourceRecord, ...]


class ImportParseError(Exception):
    """The source file cannot be parsed; the job fails. `field` names the
    offending header column where there is one; `row_number` is the source
    row the problem was found on, if any. Never carries cell values."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        row_number: int | None = None,
        field: str | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.row_number = row_number
        self.field = field


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_header(cells: Sequence[Any]) -> tuple[str, ...]:
    # Trailing blank cells are not columns (spreadsheets routinely report
    # them); a blank cell *between* named columns is a header error below.
    trimmed = list(cells)
    while trimmed and _is_blank(trimmed[-1]):
        trimmed.pop()
    if not trimmed:
        raise ImportParseError(
            "import_header_missing",
            "The first row must be a header row with canonical column names",
            row_number=_HEADER_ROW_NUMBER,
        )

    columns: list[str] = []
    for cell in trimmed:
        name = cell.strip() if isinstance(cell, str) else "" if cell is None else str(cell)
        reported = name[:_MAX_REPORTED_COLUMN_LENGTH]
        if name not in CANONICAL_IMPORT_COLUMNS:
            raise ImportParseError(
                "import_header_unknown_column",
                "The header contains a column that is not a canonical import column",
                row_number=_HEADER_ROW_NUMBER,
                field=reported or None,
            )
        if name in columns:
            raise ImportParseError(
                "import_header_duplicate_column",
                "The header contains the same column more than once",
                row_number=_HEADER_ROW_NUMBER,
                field=reported,
            )
        columns.append(name)

    for required in REQUIRED_IMPORT_COLUMNS:
        if required not in columns:
            raise ImportParseError(
                "import_header_missing_required_column",
                "The header is missing a required column",
                row_number=_HEADER_ROW_NUMBER,
                field=required,
            )
    return tuple(columns)


def _build_source(rows: Iterable[Sequence[Any]]) -> ParsedSource:
    iterator = iter(rows)
    first = next(iterator, None)
    if first is None:
        raise ImportParseError(
            "import_header_missing",
            "The file is empty; the first row must be a header row",
            row_number=_HEADER_ROW_NUMBER,
        )
    columns = _parse_header(first)

    records: list[SourceRecord] = []
    for offset, cells in enumerate(iterator):
        row_number = _HEADER_ROW_NUMBER + 1 + offset
        cells = list(cells)
        if any(not _is_blank(cell) for cell in cells[len(columns) :]):
            raise ImportParseError(
                "import_file_malformed",
                "A row has a value outside the header's columns",
                row_number=row_number,
            )
        if all(_is_blank(cell) for cell in cells):
            continue
        values = {
            column: cells[index] if index < len(cells) else None
            for index, column in enumerate(columns)
        }
        records.append(SourceRecord(row_number=row_number, values=values))
    return ParsedSource(columns=columns, records=tuple(records))


def parse_csv(content: bytes) -> ParsedSource:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportParseError(
            "import_file_unreadable", "The CSV file is not valid UTF-8 text"
        ) from exc

    if "\x00" in text:
        # Text CSV never contains NUL; Python's csv module no longer
        # rejects it by itself.
        raise ImportParseError("import_file_malformed", "The CSV file contains NUL bytes")

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=",", strict=True)
    try:
        rows = list(reader)
    except csv.Error as exc:
        # `reader.line_num` counts physical lines, which differs from the
        # record-based `row_number` once a quoted value spans lines — so it
        # is reported only in the message, never as `row_number`.
        raise ImportParseError(
            "import_file_malformed", f"The CSV file is malformed near line {reader.line_num}"
        ) from exc
    return _build_source(rows)


def parse_xlsx(content: bytes) -> ParsedSource:
    # openpyxl surfaces a corrupt/non-workbook file as any of several
    # unrelated exception types (BadZipFile, KeyError, XML parser errors,
    # ...), so every failure to open or read the first worksheet is the
    # same "unreadable" outcome for the user.
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportParseError(
            "import_file_unreadable", "The XLSX file cannot be read as a workbook"
        ) from exc

    try:
        if not workbook.worksheets:
            raise ImportParseError(
                "import_header_missing", "The workbook has no worksheet to import"
            )
        try:
            rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        except Exception as exc:
            raise ImportParseError(
                "import_file_unreadable", "The first worksheet cannot be read"
            ) from exc
    finally:
        workbook.close()
    return _build_source(rows)


def parse_source(content: bytes, source_format: str) -> ParsedSource:
    if source_format == "csv":
        return parse_csv(content)
    if source_format == "xlsx":
        return parse_xlsx(content)
    raise ValueError(f"Unsupported source_format: {source_format!r}")


__all__ = [
    "REQUIRED_IMPORT_COLUMNS",
    "OPTIONAL_IMPORT_COLUMNS",
    "CANONICAL_IMPORT_COLUMNS",
    "SourceRecord",
    "ParsedSource",
    "ImportParseError",
    "parse_csv",
    "parse_xlsx",
    "parse_source",
]
